from pathlib import Path

from docx import Document

from app.models import MatchResult, ProcessOptions, QueryPlan, Segment
from app.services.docx_exporter import export_docx


def test_exporter_creates_readable_docx(tmp_path: Path) -> None:
    segment = Segment(1, "Đây là một phân đoạn thử nghiệm cho tài liệu.", 1, 9, 4.3)
    match = MatchResult(
        segment=segment,
        query_plan=QueryPlan("business team office", [], "en-US", "business:neutral", ["business"]),
        video=None,
        selected_file=None,
        score=0.0,
    )
    output = tmp_path / "report.docx"
    export_docx([match], ProcessOptions(), "sample.txt", output)
    assert output.exists() and output.stat().st_size > 10_000
    document = Document(output)
    assert "STOCK VIDEO MATCHING PLAN" in document.paragraphs[0].text
    assert len(document.tables) == 2
