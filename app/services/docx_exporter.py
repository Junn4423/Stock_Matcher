from __future__ import annotations

import io
import logging
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import httpx
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from app.models import MatchResult, ProcessOptions

LOGGER = logging.getLogger(__name__)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_hyperlink(paragraph, text: str, url: str, bold: bool = False) -> None:
    if not url:
        paragraph.add_run(text)
        return
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _placeholder_image(text: str, size: tuple[int, int] = (640, 360)) -> bytes:
    image = Image.new("RGB", size, (235, 239, 244))
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    message = text[:48] or "No preview"
    bbox = draw.multiline_textbbox((0, 0), message, font=font, spacing=6, align="center")
    x = (size[0] - (bbox[2] - bbox[0])) / 2
    y = (size[1] - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), message, fill=(45, 55, 72), font=font, spacing=6, align="center")
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=88)
    return buffer.getvalue()


def _fetch_preview(url: str, fallback_text: str) -> bytes:
    if not url:
        return _placeholder_image(fallback_text)
    hostname = (urlparse(url).hostname or "").lower()
    if not hostname.endswith(".pexels.com"):
        return _placeholder_image(fallback_text)
    try:
        response = httpx.get(
            url,
            timeout=httpx.Timeout(15, connect=8),
            follow_redirects=True,
            headers={"User-Agent": "StockScriptMatcher/1.0"},
        )
        response.raise_for_status()
        image = Image.open(io.BytesIO(response.content)).convert("RGB")
        image.thumbnail((800, 450))
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=86, optimize=True)
        return buffer.getvalue()
    except Exception as exc:
        LOGGER.warning("Cannot load preview %s: %s", url, exc)
        return _placeholder_image(fallback_text)


def _set_document_defaults(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")


def export_docx(
    matches: list[MatchResult],
    options: ProcessOptions,
    source_name: str,
    output_path: Path,
) -> Path:
    document = Document()
    _set_document_defaults(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("STOCK VIDEO MATCHING PLAN")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(28, 49, 83)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Transcript → Visual search query → Pexels stock footage")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(10.5)
    subtitle_run.font.color.rgb = RGBColor(88, 96, 110)

    meta = document.add_table(rows=2, cols=4)
    meta.autofit = False
    meta_data = [
        ("Nguồn", source_name),
        ("Số phân đoạn", str(len(matches))),
        ("Khung hình", options.orientation),
        ("Ngày tạo", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Chế độ tách", options.segment_mode),
        ("Tốc độ đọc", f"{options.words_per_minute} từ/phút"),
        ("Truy vấn/đoạn", str(options.max_queries_per_segment)),
        ("Nguồn video", "Pexels"),
    ]
    for index, (label, value) in enumerate(meta_data):
        row = index // 4
        cell = meta.cell(row, index % 4)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(value)
        _set_cell_shading(cell, "F2F5F9" if row == 0 else "FFFFFF")
        _set_cell_margins(cell, 65, 90, 65, 90)
    document.add_paragraph()

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.autofit = False
    headers = ["#", "Kịch bản", "Ý đồ tìm kiếm", "Preview", "Video phù hợp", "Timing / Score"]
    widths = [Cm(0.7), Cm(7.0), Cm(4.5), Cm(4.0), Cm(6.7), Cm(2.8)]
    header_row = table.rows[0]
    _repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = widths[index]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "1C3153")
        _set_cell_margins(cell, 90, 90, 90, 90)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)

    for match in matches:
        row = table.add_row()
        _keep_row_together(row)
        for index, width in enumerate(widths):
            row.cells[index].width = width
            row.cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(row.cells[index], 75, 85, 75, 85)
        if match.segment.index % 2 == 0:
            for cell in row.cells:
                _set_cell_shading(cell, "F8FAFC")

        row.cells[0].text = str(match.segment.index)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].text = match.segment.text

        intent_cell = row.cells[2]
        intent_cell.text = ""
        p = intent_cell.paragraphs[0]
        run = p.add_run(match.query_plan.primary_query)
        run.bold = True
        if match.query_plan.keywords:
            p2 = intent_cell.add_paragraph("Keywords: " + ", ".join(match.query_plan.keywords))
            p2.style = document.styles["Normal"]
            for r in p2.runs:
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(92, 104, 120)
        p3 = intent_cell.add_paragraph("Intent: " + match.query_plan.visual_intent)
        for r in p3.runs:
            r.font.size = Pt(8)
            r.font.italic = True

        preview_cell = row.cells[3]
        preview_cell.text = ""
        preview_paragraph = preview_cell.paragraphs[0]
        preview_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        preview_bytes = _fetch_preview(
            match.video.image if match.video else "",
            f"Segment {match.segment.index}\nNo preview",
        )
        preview_paragraph.add_run().add_picture(io.BytesIO(preview_bytes), width=Inches(1.55))

        link_cell = row.cells[4]
        link_cell.text = ""
        link_paragraph = link_cell.paragraphs[0]
        if match.video:
            _add_hyperlink(link_paragraph, "Mở video trên Pexels", match.video.url, bold=True)
            creator_p = link_cell.add_paragraph()
            creator_p.add_run("Video by ").italic = True
            _add_hyperlink(creator_p, match.video.creator_name, match.video.creator_url)
            if match.selected_file:
                direct_p = link_cell.add_paragraph()
                _add_hyperlink(direct_p, "Direct MP4", match.selected_file.link)
                details = link_cell.add_paragraph(
                    f"{match.selected_file.width}×{match.selected_file.height} • "
                    f"{match.video.duration}s • {match.selected_file.fps:g}fps"
                )
                for r in details.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(92, 104, 120)
        else:
            missing = link_paragraph.add_run("Không tìm thấy video phù hợp")
            missing.bold = True
            missing.font.color.rgb = RGBColor(174, 45, 45)

        score_cell = row.cells[5]
        score_cell.text = ""
        timing = score_cell.paragraphs[0]
        timing.alignment = WD_ALIGN_PARAGRAPH.CENTER
        timing.add_run(f"~{match.segment.estimated_seconds:.1f}s").bold = True
        score_p = score_cell.add_paragraph(f"Score: {match.score:.2f}")
        score_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        status_p = score_cell.add_paragraph("Đã tải" if match.local_video_path else "Chưa tải")
        status_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in status_p.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(43, 122, 74) if match.local_video_path else RGBColor(92, 104, 120)

    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Videos provided by Pexels. Kiểm tra lại giấy phép và nội dung trước khi xuất bản thương mại."
    )
    note_run.font.size = Pt(8.5)
    note_run.font.color.rgb = RGBColor(92, 104, 120)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Stock Script Matcher • Pexels matching report").font.size = Pt(8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
