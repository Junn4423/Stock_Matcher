from __future__ import annotations

import re
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document
from pypdf import PdfReader


class UnsupportedFileType(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".docx", ".pdf", ".srt", ".vtt"}


def _decode_text(raw: bytes) -> str:
    best = from_bytes(raw).best()
    if best is not None:
        return str(best)
    return raw.decode("utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _strip_subtitle_markup(text: str) -> str:
    lines: list[str] = []
    for line in text.replace("\ufeff", "").splitlines():
        stripped = line.strip()
        if not stripped:
            lines.append("")
            continue
        if stripped.upper() == "WEBVTT":
            continue
        if re.fullmatch(r"\d+", stripped):
            continue
        if "-->" in stripped:
            continue
        stripped = re.sub(r"<[^>]+>", "", stripped)
        stripped = re.sub(r"\{\\[^}]+}", "", stripped)
        lines.append(stripped)
    return "\n".join(lines)


def normalize_extracted_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Định dạng {extension or '(không có phần mở rộng)'} chưa được hỗ trợ. "
            f"Hỗ trợ: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if extension == ".docx":
        text = _extract_docx(path)
    elif extension == ".pdf":
        text = _extract_pdf(path)
    else:
        text = _decode_text(path.read_bytes())
        if extension in {".srt", ".vtt"}:
            text = _strip_subtitle_markup(text)

    normalized = normalize_extracted_text(text)
    if not normalized:
        raise ValueError("Không đọc được nội dung văn bản từ file.")
    return normalized
